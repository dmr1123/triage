# -*- coding: utf-8 -*-
"""報告版（精簡）：每篇 2 頁，聚焦「訊號如何取得」＋「如何結合主訴做判斷」。
沿用 paper_content.PAPERS 的資料。"""
import sys
sys.path.insert(0, '/home/user/triage/ppt')
from deck_engine import *
from deck_engine import _blank
from paper_content import PAPERS
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = new_deck()


def _bullets(slide, l, t, w, items, size=12.5, gap=6, color=BODY, lh=1.2, h=Inches(4.2)):
    tf = textbox(slide, l, t, w, h)
    first = True
    for it in items:
        txt, lvl = (it if isinstance(it, tuple) else (it, 0))
        para(tf, txt, size - lvl, color if lvl == 0 else MUTED, bullet=True, level=lvl,
             space_after=gap, first=first, line_spacing=lh)
        first = False
    return tf


# 每篇的「結合主訴」重點 + 迷你融合流程 + 一句帶走（key = num）
REPORT = {
    "1": {
        "cc": "主訴以「結構化勾選清單」輸入（家長擔憂、呼吸困難，加觀察到的水腫、蒼白），"
              "和自動量到的生命徵象（年齡、心率、體溫、上臂圍、血氧）一起放進「同一條 Logistic 回歸公式」，"
              "各自是一個加權項，直接輸出一個住院風險機率，再用 8%／40% 兩個門檻分成紅／黃／綠。",
        "point": "重點：症狀與數字不是分開評分，而是同一條公式裡的權重——把「家長覺得不對勁」也變成一個可計算的變數。",
        "mini": (["徵象 ＋ 主訴(勾選)", "Logistic 回歸(9 變數)", "住院風險 → 紅/黃/綠"], ["Logistic 回歸(9 變數)"]),
        "take": "把主訴當成公式裡的變數，和血氧、心率等一起加權，最直接的「結合」。",
    },
    "2": {
        "cc": "主訴以勾選清單輸入（頭痛／視力、胸痛／喘、出血＋腹痛）加病史（胎次、懷孕週數），"
              "和血壓、尿蛋白、探頭測到的 SpO₂ 一起丟進既有的 miniPIERS Logistic 模型，"
              "輸出「48 小時內發生嚴重併發症」的機率，再轉成紅綠燈式檢傷。",
        "point": "重點：血氧探頭補上一個關鍵徵象——研究顯示把 SpO₂ 加進 miniPIERS，抓到高危孕婦的敏感度由 ~33% 升到 ~50%。",
        "mini": (["徵象 ＋ 症狀(勾選)", "miniPIERS 回歸", "48h 危險機率 → 分級"], ["miniPIERS 回歸"]),
        "take": "沿用成熟風險模型，手機＋血氧計把它搬到床邊；主訴與 SpO₂ 都是模型輸入。",
    },
    "3": {
        "cc": "主訴以 App 自填問卷收集，和量到的體溫／血氧／心率／咳嗽音一起「加密上傳到醫師的網頁儀表板」，"
              "由醫師（人）看著數字＋症狀趨勢判斷誰在惡化——是「人來融合」，沒有自動模型。",
        "point": "重點：這是另一條路線——不做自動評分，把客觀數據送到醫師眼前，讓專業判斷取代電話追蹤。",
        "mini": (["徵象 ＋ 自填症狀", "醫師看儀表板", "人工檢傷 / 惡化標記"], []),
        "take": "融合由「人」完成：便宜設備負責把徵象與症狀變成醫師看得到的硬數據。",
    },
    "4": {
        "cc": "主訴是護理師打的「自由文字」。模型先把文字斷詞、轉成詞向量（embedding，讓意義相近的詞數字也相近），"
              "再和生命徵象數字一起餵進神經網路的共同層，由網路自己「學」出哪些字句代表危險，輸出重症機率。",
        "point": "重點：這是最典型的「文字＋數字」自動融合。加入主訴文字後，AUC 由 0.82 升到 0.85，遠勝傳統 ESI（0.67）。",
        "mini": (["生命徵象數字 ＋ 主訴文字", "神經網路融合", "24h 重症機率"], ["神經網路融合"]),
        "take": "不需人工整理症狀，讓深度學習直接讀懂主訴文字，是最靈活的結合方式。",
    },
    "5": {
        "cc": "本篇「不使用主訴」，只做純訊號 → 血氧。它是上游的「感測積木」：量到的 SpO₂（低於 90% 警示）"
              "本身就是一個徵象，可以接到像論文 1、4 那樣的「主訴＋徵象」融合模型裡，補上「連便宜手機都能量血氧」這一塊。",
        "point": "重點：這篇解決的是『怎麼用一支普通手機取得血氧』；把它產生的 SpO₂ 餵進前幾篇的融合模型，就能結合主訴。",
        "mini": (["手機相機 PPG", "CNN 估血氧", "SpO₂(<90% 警示)"], ["CNN 估血氧"]),
        "take": "本身不含主訴，但提供最便宜的血氧來源，可當作其他融合模型的輸入。",
    },
}


def _badges_row(slide, top, dev, w=Inches(2.82), h=Inches(0.72)):
    items = [
        ("任何手機都能用？", dev["any_phone"][0], dev["any_phone"][1]),
        ("需要 App？", dev["app"][0], dev["app"][1]),
        ("開源？", dev["oss"][0], dev["oss"][1]),
        ("成本 vs NT$5000", "低成本" if dev["cost_kind"] == "ok" else ("超過" if dev["cost_kind"] == "warn" else "不適用"), dev["cost_kind"]),
    ]
    x = Inches(0.78); gap = Inches(0.14)
    for lab, val, kind in items:
        badge(slide, x, top, lab, val, kind=kind, w=w, h=h)
        x += w + gap


def _sensing_rows(slide, x, y, w, sensing, area_h=Inches(2.7)):
    n = len(sensing)
    two_col = n > 4
    if two_col:
        cols = 2
        per_col = (n + 1) // 2
        col_w = (w - Inches(0.35)) / 2
    else:
        cols = 1
        per_col = n
        col_w = w
    per_h = min(Inches(0.82), int(area_h / max(1, per_col)))
    fs = 11.5 if not two_col else 10.5
    for i, (sign, how) in enumerate(sensing):
        col = i // per_col
        row = i % per_col
        cx = x + col * (col_w + Inches(0.35))
        cy = y + row * per_h
        rect(slide, cx, cy + Inches(0.06), Inches(0.14), Inches(0.14), fill=ACCENT, shape=MSO_SHAPE.OVAL)
        tf = textbox(slide, cx + Inches(0.28), cy - Inches(0.02), col_w - Inches(0.32), per_h)
        para(tf, [(sign + "：", True, INK), (how, False, BODY)], fs, first=True, line_spacing=1.1)


def render_report_paper(prs, p):
    r = REPORT[p["num"]]
    # ---------------- Slide 1: 訊號如何取得
    s = _blank(prs)
    rect(s, 0, 0, SW, SH, fill=BG)
    rect(s, 0, 0, SW, Inches(1.35), fill=PRIMARY)
    rect(s, 0, 0, Inches(0.22), Inches(1.35), fill=ACCENT)
    tf = textbox(s, Inches(0.78), Inches(0.2), Inches(9.8), Inches(1.05))
    para(tf, p["kicker"] + "　—　① 訊號如何取得", 11.5, RGBColor(0xBF, 0xDD, 0xE7), bold=True, space_after=3, first=True)
    para(tf, p["title"], 18, WHITE, bold=True, line_spacing=1.0)
    tfn = textbox(s, Inches(11.5), Inches(0.2), Inches(1.7), Inches(1.1))
    para(tfn, p["num"] + " / 5", 30, RGBColor(0x2A, 0x6A, 0x83), bold=True, align=PP_ALIGN.RIGHT, first=True)
    # badges
    _badges_row(s, Inches(1.55), p["device"])
    # device + cost line
    tf = textbox(s, Inches(0.78), Inches(2.42), Inches(11.75), Inches(0.85))
    para(tf, [("設備：", True, PRIMARY), (p["device"]["type"], False, BODY)], 11.5, first=True, line_spacing=1.12)
    para(tf, [("金額：", True, PRIMARY), (p["device"]["cost"], False, BODY)], 11.5, line_spacing=1.12, space_before=2)
    # sensing panel (focus)
    _, sx, sy, sw = panel_titled(s, Inches(0.78), Inches(3.45), Inches(11.75), Inches(3.35), "量到哪些徵象、訊號怎麼取得", accent=ACCENT, fill=RGBColor(0xED, 0xF6, 0xF8))
    _sensing_rows(s, sx, sy, sw, p["sensing"])
    footer(s)
    notes(s, "【訊號取得重點】\n" + "\n".join(f"· {sign}：{how}" for sign, how in p["sensing"]) +
          "\n\n設備：" + p["device"]["type"] + "\n金額：" + p["device"]["cost"])

    # ---------------- Slide 2: 如何結合主訴 → 判斷
    s = _blank(prs)
    rect(s, 0, 0, SW, SH, fill=BG)
    title_bar(s, p["kicker"], "② 如何結合主訴 → 做判斷", num=None)
    # cc panel
    _, cx, cy, cw = panel_titled(s, Inches(0.78), Inches(1.55), Inches(11.75), Inches(2.15), "主訴怎麼收集、怎麼和徵象結合")
    tf = textbox(s, cx, cy, cw, Inches(1.5))
    para(tf, r["cc"], 12.5, BODY, first=True, line_spacing=1.3)
    # mini fusion flow
    labels, models = r["mini"]
    flow_snake(s, Inches(1.4), Inches(4.05), Inches(10.5), labels, models, None, box_h=Inches(1.0), max_cols=3)
    # highlight + takeaway
    hk, hv, hd = p["hl"]
    hb = rect(s, Inches(0.78), Inches(5.5), Inches(3.0), Inches(1.25), fill=PRIMARY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    tf = textbox(s, Inches(0.92), Inches(5.6), Inches(2.72), Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, hk, 11, RGBColor(0xBF, 0xDD, 0xE7), bold=True, align=PP_ALIGN.CENTER, first=True, space_after=2)
    para(tf, hv, 24, WHITE, bold=True, align=PP_ALIGN.CENTER, space_after=2)
    para(tf, hd, 9.5, RGBColor(0xCF, 0xE2, 0xEA), align=PP_ALIGN.CENTER, line_spacing=1.0)
    callout(s, Inches(4.0), Inches(5.5), Inches(8.53), Inches(1.25), "一句帶走", r["point"] + "\n" + r["take"], accent=MODEL, fill=MODELBG)
    footer(s)
    notes(s, "【結合主訴】\n" + r["cc"] + "\n\n" + r["point"] + "\n" + r["take"])


# ================================================================= TITLE
s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=PRIMD)
rect(s, 0, 0, Inches(0.24), SH, fill=ACCENT)
tf = textbox(s, Inches(0.95), Inches(2.1), Inches(11), Inches(2.8))
para(tf, "低成本設備輔助檢傷（報告版）", 38, WHITE, bold=True, space_after=8, first=True)
para(tf, "聚焦：訊號如何取得　×　如何結合主訴做判斷", 19, RGBColor(0xCF, 0xE2, 0xEA), line_spacing=1.2)
tf2 = textbox(s, Inches(0.95), Inches(5.2), Inches(11), Inches(1))
para(tf2, "5 篇論文 · 每篇 2 頁", 15, ACCENT, bold=True, first=True)

# ================================================================= 總結 ①：便宜設備能取得哪些 EHR
def _mat_pill(slide, right_x, y, text, color):
    w = Inches(est_text_w_in(text, 10.5) + 0.42)
    sp = rect(slide, right_x - w, y, w, Inches(0.36), fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, 10.5, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)

s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=BG)
title_bar(s, "總結 ①", "便宜設備現在能取得哪些 EHR 數據？", num=None)
tfi = textbox(s, Inches(0.78), Inches(1.32), Inches(11.75), Inches(0.4))
para(tfi, "每列＝ EHR 項目 ｜ 便宜取得方式 ｜ 出現於哪幾篇 ｜ 成熟度（綠=可靠　黃=篩檢/研究　紅=受限）",
     10.5, MUTED, first=True)
EAF = RGBColor(0xEA, 0xF2, 0xF6)
ehr = [
    ("心率 HR", "手機相機 PPG · 夾式血氧計 · 智慧手錶", "論文 1·2·3·5", "成熟", GREEN),
    ("血氧 SpO₂", "夾式/插入式血氧計(準) · 手機相機＋閃光(篩檢級)", "論文 1·2·3·5", "成熟 / 篩檢", GREEN),
    ("呼吸速率 RR", "App 點按(RRate) · 相機拍胸部起伏 · 麥克風", "論文 1", "尚可用", AMBER),
    ("體溫", "紅外線體溫計外掛 · 相機 OCR 讀顯示", "論文 1·3", "成熟", GREEN),
    ("血壓 BP", "壓脈帶手動輸入；純手機量測尚不可靠", "論文 2", "受限", RED),
    ("咳嗽 / 呼吸音", "手機麥克風錄音（自動分類仍研究）", "論文 3", "研究中", AMBER),
    ("臨床觀察（水腫·蒼白·意識）", "由醫護觀察後輸入 EHR 欄位", "論文 1·2", "靠人判讀", ACCENT),
]
y0 = Inches(1.86); rh = Inches(0.64)
for i, (name, method, papers, mat, color) in enumerate(ehr):
    y = y0 + i * rh
    rect(s, Inches(0.78), y, Inches(11.75), Inches(0.54), fill=(PANEL if i % 2 else EAF), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, Inches(0.78), y, Inches(0.09), Inches(0.54), fill=color)
    tfn = textbox(s, Inches(1.0), y + Inches(0.05), Inches(2.95), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfn, name, 12, INK, bold=True, first=True, line_spacing=1.0)
    tfm = textbox(s, Inches(4.0), y + Inches(0.05), Inches(4.9), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfm, method, 10.5, BODY, first=True, line_spacing=1.0)
    tfp = textbox(s, Inches(8.98), y + Inches(0.05), Inches(1.95), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfp, papers, 10.5, ACCENT, bold=True, first=True, line_spacing=1.0)
    _mat_pill(s, Inches(12.42), y + Inches(0.09), mat, color)
callout(s, Inches(0.78), Inches(6.42), Inches(11.75), Inches(0.6), "",
        "註：論文 4（Joseph）不新增感測，直接沿用 EHR 既有生命徵象數字；上表列的是「用便宜設備主動取得」的對應。"
        "真正穩定 <NT$5000 的以心率、血氧、體溫、呼吸為主，血壓與聲音類仍是弱點。", accent=PRIMARY, fill=PANEL)

# ================================================================= 總結 ①b：只用相機能做哪些
s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=BG)
title_bar(s, "總結 ①b · 聚焦相機", "只用手機相機（含閃光），能做到哪些？", num=None)
tfi = textbox(s, Inches(0.78), Inches(1.32), Inches(11.75), Inches(0.4))
para(tfi, "把上一頁篩成「純相機、不外接任何硬體」能做的。狀態：綠=可靠　黃=篩檢/研究　紅=做不到/實驗", 10.5, MUTED, first=True)
cam = [
    ("心率 HR", "手指蓋後鏡頭＋閃光 → 接觸式 PPG（血液脈動亮暗）", "論文 3·5", "可靠", GREEN),
    ("血氧 SpO₂", "同段指尖影片，比紅/藍綠光脈動比值 → CNN 估算", "論文 5", "篩檢級", AMBER),
    ("讀螢幕數字（OCR）", "拍血氧計/體溫計螢幕 → OCR 讀數字（間接）", "論文 3", "間接·成熟", GREEN),
    ("呼吸速率 RR", "指尖 PPG 基線起伏，或前鏡頭拍胸部起伏", "一般方法", "可行·研究", AMBER),
    ("體溫", "相機測不到熱 → 不能直測，只能 OCR 讀體溫計", "論文 3 (OCR)", "不能直測", RED),
    ("血壓 BP", "PPG 波形估血壓仍實驗、不可靠、非臨床級", "非這 5 篇", "實驗", RED),
    ("咳嗽 / 呼吸音", "屬手機「麥克風」，不是相機的工作", "—", "非相機", MUTED),
]
y0 = Inches(1.86); rh = Inches(0.64)
for i, (name, how, papers, st, color) in enumerate(cam):
    y = y0 + i * rh
    rect(s, Inches(0.78), y, Inches(11.75), Inches(0.54), fill=(PANEL if i % 2 else EAF), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, Inches(0.78), y, Inches(0.09), Inches(0.54), fill=color)
    tfn = textbox(s, Inches(1.0), y + Inches(0.05), Inches(2.85), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfn, name, 12, INK, bold=True, first=True, line_spacing=1.0)
    tfm = textbox(s, Inches(3.9), y + Inches(0.05), Inches(5.0), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfm, how, 10.5, BODY, first=True, line_spacing=1.0)
    tfp = textbox(s, Inches(8.98), y + Inches(0.05), Inches(1.95), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
    para(tfp, papers, 10.5, ACCENT, bold=True, first=True, line_spacing=1.0)
    _mat_pill(s, Inches(12.42), y + Inches(0.09), st, color)
callout(s, Inches(0.78), Inches(6.34), Inches(11.75), Inches(0.66), "核心",
        "相機其實只做兩件事：① 當「接觸式血氧感測器」（指尖 PPG → 心率可靠、血氧篩檢級）；② 當「掃描器」（OCR 讀既有裝置螢幕）。體溫、血壓相機都做不好。",
        accent=MODEL, fill=MODELBG)

# ================================================================= 總結 ②：主訴 × EHR 整合方法
s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=BG)
title_bar(s, "總結 ②", "現在「主訴 × EHR」的整合方法", num=None)
methods = [
    ("A. 統計 / 機器學習模型", MODEL,
     "把症狀編成 0/1，和 EHR 數字放進同一條公式加權計算。",
     "→ 論文 1 Smart Triage、論文 2 PIERS（logistic 回歸）",
     "優：簡單、可解釋　　缺：症狀需先設計成勾選項"),
    ("B. 深度學習讀「自由文字」", MODEL,
     "主訴文字→詞向量(embedding)→神經網路，與 EHR 數字融合。",
     "→ 論文 4 Joseph（加主訴 AUC 0.82→0.85）",
     "優：最靈活、免人工整理　　缺：黑盒、需大量資料"),
    ("C. 人工融合（醫師）", ACCENT,
     "EHR 數據＋症狀上傳儀表板，由醫師看趨勢判讀。",
     "→ 論文 3 e-CoVig",
     "優：免訓練資料、可解釋　　缺：仍要人力"),
    ("D. 規則 / 危險徵象清單（傳統）", MUTED,
     "用固定規則，如 WHO 危險徵象、ESI 分級。",
     "→ 論文 4 用 ESI 當對照；論文 1 勾選項源自 WHO 危險徵象",
     "優：透明、好推廣　　缺：較不準（ESI AUC≈0.67）"),
]
for i, (title, accent, how, ex, pc) in enumerate(methods):
    col = i % 2; row = i // 2
    x = Inches(0.78) + col * Inches(6.0)
    y = Inches(1.65) + row * Inches(2.5)
    rect(s, x, y, Inches(5.75), Inches(2.3), fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x, y, Inches(0.11), Inches(2.3), fill=accent)
    tf = textbox(s, x + Inches(0.3), y + Inches(0.18), Inches(5.25), Inches(2.0))
    para(tf, title, 14, PRIMARY, bold=True, space_after=6, first=True)
    para(tf, how, 11.5, BODY, space_after=6, line_spacing=1.16)
    para(tf, ex, 11, ACCENT, bold=True, space_after=5, line_spacing=1.12)
    para(tf, pc, 10.5, MUTED, line_spacing=1.12)
callout(s, Inches(0.78), Inches(6.55), Inches(11.75), Inches(0.44), "",
        "四種方法可混用：多數系統用 A/B 自動算風險，或用 C 讓醫師把關；A、B 正是「把主訴變成模型輸入」。", accent=PRIMARY, fill=PANEL)

# ================================================================= OVERVIEW
s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=BG)
title_bar(s, "總覽", "共通主軸：取得訊號 → 結合主訴 → 判斷", num=None)
# recipe flow
flow_snake(s, Inches(1.4), Inches(1.7), Inches(10.5), ["便宜設備取得徵象", "結合主訴/症狀", "判斷（模型或醫師）"],
           ["判斷（模型或醫師）"], {"便宜設備取得徵象": "手機相機·血氧計·麥克風·外掛", "結合主訴/症狀": "自由文字 或 勾選清單", "判斷（模型或醫師）": "檢傷等級·風險機率"},
           box_h=Inches(1.05), max_cols=3)
headers = ["#　論文", "訊號怎麼取得", "主訴怎麼結合"]
col_w = [Inches(2.35), Inches(4.7), Inches(4.7)]
rows = [
    ["1　Smart Triage", "夾式血氧計讀 SpO₂/心率＋人工量測", "勾選症狀與徵象同進 Logistic 公式"],
    ["2　PIERS", "手機接血氧探頭讀 SpO₂＋手動輸入", "症狀勾選＋SpO₂ 進 miniPIERS 模型"],
    ["3　e-CoVig", "DIY 外掛/相機測體溫·血氧·咳嗽音", "自填症狀＋數據上傳，醫師人工判"],
    ["4　Joseph", "沿用急診既有生命徵象數字", "主訴自由文字→詞向量，神經網路融合"],
    ["5　Hoffman", "手指按相機＋閃光→PPG→CNN 估血氧", "本篇不用主訴；SpO₂ 可接入他人模型"],
]
simple_table(s, Inches(0.78), Inches(3.35), col_w, headers, rows, row_h=Inches(0.62), fs=11)
footer(s)

# ================================================================= PAPERS
for p in PAPERS:
    render_report_paper(prs, p)

# ================================================================= CLOSING SUMMARY
s = _blank(prs)
rect(s, 0, 0, SW, SH, fill=BG)
title_bar(s, "小結", "兩種「結合主訴」的路線", num=None)
_, ax, ay, aw = panel_titled(s, Inches(0.78), Inches(1.6), Inches(5.75), Inches(4.4), "A. 自動融合（模型）", accent=MODEL, fill=MODELBG)
_bullets(s, ax, ay, aw, [
    "把徵象數字＋主訴一起餵給模型，直接輸出風險/分級",
    "論文 1 Smart Triage：勾選症狀＋SpO₂ 進 Logistic",
    "論文 2 PIERS：症狀＋SpO₂ 進 miniPIERS",
    "論文 4 Joseph：主訴『自由文字』進神經網路（最靈活）",
    "論文 5 Hoffman：提供便宜 SpO₂，可當上游輸入",
], size=12, gap=9)
_, bx, by, bw = panel_titled(s, Inches(6.78), Inches(1.6), Inches(5.75), Inches(4.4), "B. 人工融合（醫師）", accent=ACCENT)
_bullets(s, bx, by, bw, [
    "便宜設備只負責把徵象＋症狀變成客觀數據",
    "上傳儀表板，由醫師看趨勢判斷惡化",
    "論文 3 e-CoVig 走這條路",
    "優點：不需訓練資料、可解釋；缺點：仍要人力",
], size=12, gap=9)
callout(s, Inches(0.78), Inches(6.2), Inches(11.75), Inches(0.62), "",
        "訊號取得越便宜（手機相機/麥克風/外掛），就越能把「量徵象＋聽主訴」下放到第一線，補上沒有專業設備的缺口。", accent=PRIMARY, fill=PANEL)
footer(s)

OUT = '/home/user/triage/ppt/低成本設備輔助檢傷_報告版.pptx'
prs.save(OUT)
print("report deck saved:", len(prs.slides._sldIdLst), "slides ->", OUT)
