# -*- coding: utf-8 -*-
"""Render a .pptx to per-slide PNG wireframes with Pillow (for layout QA only).
Approximates shapes, fills, and text; not a faithful renderer but good enough to
catch overflow, overlap, and alignment problems."""
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_IN = 914400
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"

def load_font(px):
    px = max(6, int(px))
    try:
        return ImageFont.truetype(FONT_PATH, px, index=0)
    except Exception:
        return ImageFont.load_default()

FCACHE = {}
def font(px):
    px = max(6, int(px))
    if px not in FCACHE:
        FCACHE[px] = load_font(px)
    return FCACHE[px]

def rgb(color):
    try:
        return (color[0], color[1], color[2])
    except Exception:
        return None

def shape_fill(sh):
    try:
        f = sh.fill
        if f.type is not None and f.type == 1:  # solid
            c = f.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None

def shape_line(sh):
    try:
        c = sh.line.color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return None

def run_color(run):
    try:
        c = run.font.color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return (40, 50, 60)

def wrap_text(draw, text, fnt, max_w):
    # char-based wrap (works for CJK + latin mix)
    lines = []
    cur = ""
    for ch in text:
        if ch == "\n":
            lines.append(cur); cur = ""; continue
        test = cur + ch
        w = draw.textlength(test, font=fnt)
        if w > max_w and cur:
            lines.append(cur); cur = ch
        else:
            cur = test
    if cur:
        lines.append(cur)
    return lines

def render(pptx_path, out_prefix, scale=96):
    prs = Presentation(pptx_path)
    sw = int(prs.slide_width / EMU_PER_IN * scale)
    sh = int(prs.slide_height / EMU_PER_IN * scale)
    paths = []
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (sw, sh), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for shp in slide.shapes:
            try:
                l = int(shp.left / EMU_PER_IN * scale)
                t = int(shp.top / EMU_PER_IN * scale)
                w = int(shp.width / EMU_PER_IN * scale)
                h = int(shp.height / EMU_PER_IN * scale)
            except Exception:
                continue
            fill = shape_fill(shp)
            line = shape_line(shp)
            is_text_only = shp.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            if (fill or line) and not is_text_only:
                d.rectangle([l, t, l + w, t + h], fill=fill, outline=line, width=2 if line else 0)
            # text
            if shp.has_text_frame:
                tf = shp.text_frame
                # approximate vertical anchor
                y = t + int(0.04 * scale)
                for p in tf.paragraphs:
                    txt = "".join(r.text for r in p.runs) or p.text
                    if not txt:
                        y += int(scale * 0.12); continue
                    sz = None
                    col = (40, 50, 60)
                    for r in p.runs:
                        if r.font.size:
                            sz = r.font.size.pt
                        col = run_color(r)
                        break
                    if sz is None:
                        sz = 14
                    fpx = sz * scale / 72.0
                    fnt = font(fpx)
                    align = p.alignment
                    max_w = w - int(0.12 * scale)
                    for ln in wrap_text(d, txt, fnt, max_w):
                        lw = d.textlength(ln, font=fnt)
                        if str(align) == "CENTER (2)" or (align is not None and int(align) == 2):
                            x = l + (w - lw) / 2
                        elif align is not None and int(align) == 3:  # right
                            x = l + w - lw - int(0.06*scale)
                        else:
                            x = l + int(0.06 * scale)
                        d.text((x, y), ln, fill=col, font=fnt)
                        y += int(fpx * 1.25)
        # slide border for reference
        d.rectangle([0, 0, sw - 1, sh - 1], outline=(200, 200, 200), width=1)
        p = f"{out_prefix}_{idx:02d}.png"
        img.save(p)
        paths.append(p)
    return paths

if __name__ == "__main__":
    pptx = sys.argv[1] if len(sys.argv) > 1 else "_test.pptx"
    pref = sys.argv[2] if len(sys.argv) > 2 else "prev"
    ps = render(pptx, pref)
    print("rendered", len(ps), "slides:", ", ".join(ps))
